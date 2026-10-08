"""Unit tests for load_india_hourly and _mask_implausible in terra.real.loaders."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from terra.real.loaders import _mask_implausible, load_india_hourly


def _create_layout_a_wb(path: Path) -> None:
    """Create a minimal openpyxl workbook mimicking Layout A (5-minute data with tag header).

    Contains 12 5-minute intervals covering (09:00, 10:00] IST (09:05 through 10:00).
    """
    wb = openpyxl.Workbook()
    # Sheet 1: _osi_config (internal junk)
    ws_osi = wb.active
    ws_osi.title = "_osi_config"
    ws_osi.append(['{"r": 300.0, "tc": []}'])

    # Sheet 2: Sheet1 (actual data)
    ws_data = wb.create_sheet(title="Sheet1")
    ws_data.append(
        ["Time", "SCADA/ANALOG/044MQ067/0", "SCADA/ANALOG/044MQ206/0", "SCADA/ANALOG/044MQ070/0"]
    )
    # Row 0 of data holds SCADA tag names
    ws_data.append([None, "NLDC_DEMAND|P", "ALL_IND_SOLAR|P", "ALL_INDIA_WIND|P"])
    # 12 rows of 5-minute intervals covering (09:00, 10:00] IST
    for m in range(5, 65, 5):
        hh = "09" if m < 60 else "10"
        mm = f"{m:02d}" if m < 60 else "00"
        ws_data.append([f"2022-01-01 {hh}:{mm}:00", 160000.0, 500.0, 5000.0])

    # Sheet 3: Sheet2 (side sheet)
    ws_side = wb.create_sheet(title="Sheet2")
    ws_side.append(["Time", "Unnamed: 1", "Unnamed: 2"])
    ws_side.append([None, "Solar+Wind", "Every Hourly Timestamps"])
    ws_side.append(["2022-01-01 10:00:00", 5000.0, 5000.0])

    wb.save(path)
    wb.close()


def _create_layout_b_wb(path: Path) -> None:
    """Create a minimal openpyxl workbook mimicking Layout B (hourly data with summary rows)."""
    wb = openpyxl.Workbook()
    ws_rep = wb.active
    ws_rep.title = "Report"
    ws_rep.append(["Timestamp", "Demand (MW)", "Wind (MW)", "Solar (MW)", "Total Generation (MW)"])
    ws_rep.append(["01-01-2024 10:00:00", 180000.0, 6000.0, 25000.0, 211000.0])
    ws_rep.append(["01-01-2024 11:00:00", 185000.0, 6200.0, 30000.0, 221200.0])
    # Trailing summary junk rows
    ws_rep.append(["Average", 182500.0, 6100.0, 27500.0, 216100.0])
    ws_rep.append(["Sum", 365000.0, 12200.0, 55000.0, 432200.0])

    wb.create_sheet(title="Sheet2")
    wb.create_sheet(title="Sheet3")

    wb.save(path)
    wb.close()


def test_load_india_hourly_layout_a(tmp_path: Path) -> None:
    """Test Layout A parsing: sheet selection, tag resolution, hour-ending resampling, IST->UTC shift."""
    wb_path = tmp_path / "September_2021.xlsx"
    _create_layout_a_wb(wb_path)

    df = load_india_hourly(tmp_path)

    assert list(df.columns) == ["demand_mw", "solar_mw", "wind_mw"]
    assert df.index.name == "ts_utc"
    assert str(df.index.tz) == "UTC"
    assert len(df) == 1  # 12 5-minute rows in (09:00, 10:00] IST aggregated to 1 hourly row labelled 10:00 IST

    # 10:00 IST = 04:30 UTC (hour-ending label, exactly 5h30m shift)
    expected_utc = pd.Timestamp("2022-01-01 04:30:00", tz="UTC")
    assert df.index[0] == expected_utc
    shift = pd.Timestamp("2022-01-01 10:00:00") - df.index[0].tz_localize(None)
    assert shift == pd.Timedelta(hours=5, minutes=30)

    # Values and dtypes
    assert df["demand_mw"].iloc[0] == pytest.approx(160000.0)
    assert df["solar_mw"].iloc[0] == pytest.approx(500.0)
    assert df["wind_mw"].iloc[0] == pytest.approx(5000.0)
    assert pd.api.types.is_float_dtype(df["demand_mw"])
    assert pd.api.types.is_float_dtype(df["solar_mw"])
    assert pd.api.types.is_float_dtype(df["wind_mw"])


def test_load_india_hourly_layout_b(tmp_path: Path) -> None:
    """Test Layout B parsing: header mapping, dropping unparseable rows, IST->UTC shift."""
    wb_path = tmp_path / "January_2024-_June_2025.xlsx"
    _create_layout_b_wb(wb_path)

    df = load_india_hourly(tmp_path)

    assert list(df.columns) == ["demand_mw", "solar_mw", "wind_mw"]
    assert df.index.name == "ts_utc"
    assert str(df.index.tz) == "UTC"
    assert len(df) == 2  # Junk 'Average' and 'Sum' rows dropped

    # IST 01-01-2024 10:00:00 is UTC 2024-01-01 04:30:00 (exactly 5h30m shift)
    expected_utc = pd.Timestamp("2024-01-01 04:30:00", tz="UTC")
    assert df.index[0] == expected_utc
    shift = pd.Timestamp("2024-01-01 10:00:00") - df.index[0].tz_localize(None)
    assert shift == pd.Timedelta(hours=5, minutes=30)

    assert df["demand_mw"].iloc[0] == pytest.approx(180000.0)
    assert df["solar_mw"].iloc[0] == pytest.approx(25000.0)
    assert df["wind_mw"].iloc[0] == pytest.approx(6000.0)
    assert pd.api.types.is_float_dtype(df["demand_mw"])
    assert pd.api.types.is_float_dtype(df["solar_mw"])
    assert pd.api.types.is_float_dtype(df["wind_mw"])


def test_load_india_hourly_combined_sorted_unique(tmp_path: Path) -> None:
    """Test multiple files are combined, sorted chronologically, and deduplicated."""
    _create_layout_a_wb(tmp_path / "file1.xlsx")
    _create_layout_b_wb(tmp_path / "file2.xlsx")

    df = load_india_hourly(tmp_path)

    assert len(df) == 3
    assert df.index.is_monotonic_increasing
    assert not df.index.has_duplicates


def test_load_india_hourly_junk_file(tmp_path: Path) -> None:
    """Test ValueError is raised with the filename if layout is unrecognized."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ColA", "ColB", "ColC"])
    ws.append([1, 2, 3])
    junk_path = tmp_path / "junk.xlsx"
    wb.save(junk_path)
    wb.close()

    with pytest.raises(ValueError, match="Unrecognized layout in file: junk.xlsx"):
        load_india_hourly(tmp_path)


def test_load_india_hourly_empty_dir(tmp_path: Path) -> None:
    """Test FileNotFoundError is raised when directory has no matching files."""
    with pytest.raises(FileNotFoundError, match="no files in"):
        load_india_hourly(tmp_path)


def test_mask_implausible_demand_rules() -> None:
    """Demand rule: demand <= 0 or demand < 0.5 * monthly median -> NaN; normal values untouched."""
    times = pd.date_range("2022-01-01 04:30", periods=6, freq="h", tz="UTC")
    raw = pd.DataFrame(
        {
            "demand_mw": [180000.0, 180000.0, 180000.0, 0.0, -500.0, 40000.0],
            "solar_mw": [1000.0] * 6,
            "wind_mw": [5000.0] * 6,
        },
        index=times,
    )
    cleaned = _mask_implausible(raw)

    # Normal values untouched
    assert cleaned["demand_mw"].iloc[0] == pytest.approx(180000.0)
    assert cleaned["demand_mw"].iloc[1] == pytest.approx(180000.0)
    assert cleaned["demand_mw"].iloc[2] == pytest.approx(180000.0)

    # <= 0 masked
    assert pd.isna(cleaned["demand_mw"].iloc[3])
    assert pd.isna(cleaned["demand_mw"].iloc[4])

    # < 0.5 * monthly median masked
    assert pd.isna(cleaned["demand_mw"].iloc[5])


def test_mask_implausible_solar_tare_clipping() -> None:
    """Solar tare rule: negative values clipped to 0.0, non-negative values untouched."""
    times = pd.date_range("2022-01-01 04:30", periods=4, freq="h", tz="UTC")
    raw = pd.DataFrame(
        {
            "demand_mw": [160000.0, 160000.0, 160000.0, 160000.0],
            "solar_mw": [-69.5, -0.01, 0.0, 850.0],
            "wind_mw": [3000.0, 3000.0, 3000.0, 3000.0],
        },
        index=times,
    )
    cleaned = _mask_implausible(raw)

    assert cleaned["solar_mw"].iloc[0] == 0.0
    assert cleaned["solar_mw"].iloc[1] == 0.0
    assert cleaned["solar_mw"].iloc[2] == 0.0
    assert cleaned["solar_mw"].iloc[3] == pytest.approx(850.0)


def test_mask_implausible_night_solar_spikes() -> None:
    """Night solar rule: IST hours 21..4 with solar > 1500 -> NaN; <= 1500 or daytime > 1500 untouched."""
    # IST 21:00 is UTC 15:30; IST 23:00 is UTC 17:30; IST 04:00 is UTC 22:30; IST 05:00 is UTC 23:30; IST 12:00 is UTC 06:30
    times = pd.to_datetime(
        [
            "2022-01-01 15:30:00+00:00",  # IST 21:00 (night)
            "2022-01-01 17:30:00+00:00",  # IST 23:00 (night)
            "2022-01-01 22:30:00+00:00",  # IST 04:00 (night)
            "2022-01-01 23:30:00+00:00",  # IST 05:00 (dawn/day, hour 5)
            "2022-01-02 06:30:00+00:00",  # IST 12:00 (day, hour 12)
        ]
    )
    raw = pd.DataFrame(
        {
            "demand_mw": [160000.0] * 5,
            "solar_mw": [5106.7, 250.0, 1800.0, 2000.0, 35000.0],
            "wind_mw": [4000.0] * 5,
        },
        index=times,
    )
    cleaned = _mask_implausible(raw)

    # 21:00 IST with 5106.7 MW (> 1500) -> masked to NaN
    assert pd.isna(cleaned["solar_mw"].iloc[0])
    # 23:00 IST with 250.0 MW (<= 1500) -> untouched
    assert cleaned["solar_mw"].iloc[1] == pytest.approx(250.0)
    # 04:00 IST with 1800.0 MW (> 1500) -> masked to NaN
    assert pd.isna(cleaned["solar_mw"].iloc[2])
    # 05:00 IST with 2000.0 MW (hour 5 not in 21..4) -> untouched
    assert cleaned["solar_mw"].iloc[3] == pytest.approx(2000.0)
    # 12:00 IST with 35000.0 MW (daytime peak) -> untouched
    assert cleaned["solar_mw"].iloc[4] == pytest.approx(35000.0)


def test_mask_implausible_wind_untouched_and_nans_preserved() -> None:
    """Wind is untouched and existing NaNs are not filled or interpolated."""
    times = pd.date_range("2022-01-01 04:30", periods=3, freq="h", tz="UTC")
    raw = pd.DataFrame(
        {
            "demand_mw": [160000.0, float("nan"), 170000.0],
            "solar_mw": [float("nan"), 500.0, float("nan")],
            "wind_mw": [5000.0, -100.0, 0.0],
        },
        index=times,
    )
    cleaned = _mask_implausible(raw)

    # Wind untouched
    assert cleaned["wind_mw"].iloc[0] == pytest.approx(5000.0)
    assert cleaned["wind_mw"].iloc[1] == pytest.approx(-100.0)
    assert cleaned["wind_mw"].iloc[2] == pytest.approx(0.0)

    # Existing NaNs remain NaN
    assert pd.isna(cleaned["demand_mw"].iloc[1])
    assert pd.isna(cleaned["solar_mw"].iloc[0])
    assert pd.isna(cleaned["solar_mw"].iloc[2])
