"""One-page jury slide: what the AI takes in, gives out, and how the system uses it.

Figures come from jury-evidence/tables/model_comparison_test.csv (written by make_jury_evidence.py).
Needs python-pptx (not a project dependency): uv pip install --python .venv/bin/python python-pptx
"""
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
INK, MUTED, PAPER, SURF, ACC, LINE = C(0x17, 0x25, 0x1E), C(0x58, 0x6A, 0x60), C(0xED, 0xF0, 0xE8), C(0xFA, 0xFB, 0xF7), C(0x1E, 0x5A, 0x43), C(0xD5, 0xDC, 0xCE)

t = pd.read_csv(ROOT / "jury-evidence/tables/model_comparison_test.csv")
sk = lambda s, v: 100 * float(t[(t.source == s) & (t.view == v) & (t.model == "ensemble")]["skill_vs_persistence"].iloc[0])
solar_skill, wind_skill = sk("solar", "daylight hours"), sk("wind", "all hours")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER


def text(x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment, para.space_after = align, Pt(5)
        r = para.add_run(); r.text = p; r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"
    return tb


def card(x, y, w, h, title, lines, head=ACC):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = 0.04; b.fill.solid(); b.fill.fore_color.rgb = SURF; b.line.color.rgb = LINE; b.shadow.inherit = False
    text(x + .2, y + .15, w - .4, .4, title, 17, head, True)
    text(x + .2, y + .65, w - .4, h - .8, lines, 12.5, INK)


text(.6, .35, 12, .7, "What the AI takes in, gives back, and how we use it", 30, INK, True)
text(.6, 1.05, 12, .4, "Hourly solar and wind generation for the next 24 to 48 hours, with an honest range around every number.", 15, MUTED)

W, H, Y = 2.85, 4.2, 1.65
card(.6, Y, W, H, "1  What goes in", [
    "Tomorrow's weather forecast: sunshine, cloud, wind speed and direction, temperature, humidity, pressure, rain.",
    "Sun position and the calendar: hour, season, daylight.",
    "What the plant produced up to now, never after.",
    "How many hours ahead we are predicting.",
    "No actual weather is ever used as input."])
card(.6 + (W + .25), Y, W, H, "2  Physics + AI", [
    "Physics model: what a solar or wind plant should produce in this weather.",
    "LightGBM (the AI): learns from history where the physics model is wrong and corrects it.",
    "A weighted blend of the two, with weights set separately for near and far hours.",
    "Ranges are then calibrated on a separate window so they hold up."])
card(.6 + 2 * (W + .25), Y, W, H, "3  What comes out", [
    "For every hour, a range in MW instead of one number:",
    "Median: the best guess",
    "P10 to P90: 80% chance the real value lands inside",
    "P5 to P95: the outer edges",
    "A trust score from 0 to 100 for each hour."])
card(.6 + 3 * (W + .25), Y, W, H, "4  How we use it", [
    "Alerts for low or high generation.",
    "Battery and backup plan for the day.",
    "Deviation-charge estimate (illustrative rates).",
    "Trust score marks which hours to rely on.",
    "Dashboard shows forecast against what happened."])

for i in range(3):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(.6 + (i + 1) * (W + .25) - .24), Inches(Y + H / 2 - .14), Inches(.23), Inches(.28))
    a.fill.solid(); a.fill.fore_color.rgb = ACC; a.line.fill.background()

bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.6), Inches(6.1), Inches(12.1), Inches(.95))
bar.adjustments[0] = 0.12; bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background(); bar.shadow.inherit = False
text(.85, 6.15, 11.6, .85, [
    f"On a test period the model never saw, the blend makes about {solar_skill:.0f}% less error than the persistence baseline for solar (daylight) and {wind_skill:.0f}% less for wind.",
    "Real Open-Meteo weather, virtual digital-twin plant at a real location. All figures come from our evaluation harness."],
    14, C(255, 255, 255), False, anchor=MSO_ANCHOR.MIDDLE)

out = ROOT / "jury-evidence/AI_one_page_slide.pptx"
prs.save(out); print(out)
